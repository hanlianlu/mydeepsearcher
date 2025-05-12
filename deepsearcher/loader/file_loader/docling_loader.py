import os
import re
import json
import logging
import tempfile
from typing import Any, Dict, List, Optional, Union

import torch
import numpy as np
import cv2
from PIL import Image
import easyocr
import pandas as pd

from langchain_core.documents import Document
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions, EasyOcrOptions
from docling.chunking import HierarchicalChunker

from deepsearcher.loader.file_loader.base import BaseLoader

logger = logging.getLogger(__name__)

class DoclingLoader(BaseLoader):
    def __init__(
        self,
        chunk_size: int = 3000,
        enable_smart_ocr: bool = True,
        ocr_config: Optional[Dict[str, Union[str, List[str]]]] = None,
        ocr_threshold: int = 200,
        max_image_size: int = 2560,
        max_table_rows: int = 500,
        max_sheet_rows: int = 2000,
        max_sheet_cols: int = 2000,
        **kwargs,
    ):
        super().__init__(**kwargs)
        self.chunk_size = chunk_size
        self.enable_smart_ocr = enable_smart_ocr
        self.ocr_threshold = ocr_threshold
        self.max_image_size = max_image_size
        self.max_table_rows = max_table_rows
        self.max_sheet_rows = max_sheet_rows
        self.max_sheet_cols = max_sheet_cols

        # OCR configuration
        self.ocr_config = ocr_config or {"engine": "easyocr", "languages": ["en", "ch_sim"]}

        # Docling converters
        self.default_converter = DocumentConverter()
        ocr_opts = EasyOcrOptions(lang=self.ocr_config["languages"])
        pdf_opts = PdfPipelineOptions(); pdf_opts.ocr_options = ocr_opts
        self.ocr_converter = DocumentConverter(
            format_options={ InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_opts) }
        )

        # Chunker
        self.chunker = HierarchicalChunker()

        # Supported file types
        self._image_exts = {"png", "jpg", "jpeg", "bmp", "tif", "tiff"}
        self._excel_exts = {"xlsx", "xls"}

        # EasyOCR reader
        try:
            self.reader = easyocr.Reader(self.ocr_config["languages"], gpu=torch.cuda.is_available())
        except Exception as e:
            self.reader = None
            logger.warning("EasyOCR init failed: %s", e)

    def _sanitize_metadata(self, meta: Dict[str, Any]) -> Dict[str, Any]:
        """
        Ensure metadata JSON under 60KB by dropping or truncating non-essential fields.
        """
        max_bytes = 60_000
        try:
            meta_json = json.dumps(meta, ensure_ascii=False)
        except Exception:
            return meta
        if len(meta_json) <= max_bytes:
            return meta
        # Drop less-critical keys
        for key in ("section", "vlm_caption", "column_headers"):
            if key in meta:
                del meta[key]
                meta_json = json.dumps(meta, ensure_ascii=False)
                if len(meta_json) <= max_bytes:
                    return meta
        # As last resort, keep only essentials
        essentials = {k: meta[k] for k in ("reference", "file_name", "type") if k in meta}
        return essentials

    def _get_process_path(self, file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower().lstrip(".")
        if ext not in self._image_exts:
            return file_path
        try:
            img = Image.open(file_path)
            w, h = img.size
            if w > self.max_image_size or h > self.max_image_size:
                scale = min(self.max_image_size / w, self.max_image_size / h)
                nw, nh = int(w*scale), int(h*scale)
                arr = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
                small = cv2.resize(arr, (nw, nh), interpolation=cv2.INTER_LANCZOS4)
                small_img = Image.fromarray(cv2.cvtColor(small, cv2.COLOR_BGR2RGB))
                with tempfile.NamedTemporaryFile(delete=True, suffix=f".{ext}") as tf:
                    small_img.save(tf.name)
                    tw, th = Image.open(tf.name).size
                    if tw > 32767 or th > 32767:
                        raise ValueError("Resized image still too large")
                    return tf.name
            return file_path
        except Exception as e:
            logger.error("Image processing failed for %s: %s", file_path, e)
            raise

    def _get_page_text(self, page: Any) -> str:
        text = ""
        try:
            for elem in getattr(page, "elements", []):
                if hasattr(elem, "text"):
                    text += elem.text + " "
        except Exception as e:
            logger.warning("Page text extract failed: %s", e)
        return text.strip()

    def load_file(self, file_path: str) -> List[Document]:
        if not os.path.exists(file_path):
            logger.error("Not found: %s", file_path)
            raise FileNotFoundError(file_path)

        ext = os.path.splitext(file_path)[1].lower().lstrip(".")
        proc = self._get_process_path(file_path)

        # Excel branch with size-check and fallback
        if ext in self._excel_exts:
            oversized = False
            try:
                import openpyxl
                wb = openpyxl.load_workbook(proc, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    dim = ws.calculate_dimension()
                    _, max_cell = dim.split(":")
                    col_letters = "".join(filter(str.isalpha, max_cell))
                    row_digits = "".join(filter(str.isdigit, max_cell))
                    max_row = int(row_digits)
                    max_col = sum(
                        (ord(c.upper()) - ord("A") + 1) * (26**i)
                        for i, c in enumerate(reversed(col_letters))
                    )
                    if max_row > self.max_sheet_rows or max_col > self.max_sheet_cols:
                        oversized = True
                        logger.warning(
                            "Sheet '%s' is %dx%d > %dx%d, falling back to pandas",
                            ws.title, max_row, max_col,
                            self.max_sheet_rows, self.max_sheet_cols
                        )
                        break
            except ImportError:
                logger.debug("openpyxl not installed; skipping size-check")
            except Exception as e:
                logger.warning("Excel size-check failed: %s", e)

            if oversized:
                return self._process_excel_with_pandas(file_path)

            # native Docling Excel ingestion
            logger.info("Docling Excel ingest: %s", file_path)
            doc = self._convert(self.default_converter, proc)
            used_ocr = False

        # PDF with smart OCR
        elif ext == "pdf" and self.enable_smart_ocr:
            doc = self._convert(self.default_converter, proc)
            used_ocr = False
            for p in getattr(doc, "pages", []):
                if len(self._get_page_text(p)) < self.ocr_threshold:
                    logger.info("Triggering PDF OCR on %s", file_path)
                    doc = self._convert(self.ocr_converter, proc)
                    used_ocr = True
                    break

        # Image with OCR detection
        elif ext in self._image_exts and self.enable_smart_ocr and self.reader:
            used_ocr = False
            try:
                img = np.array(Image.open(proc).convert("RGB"))
                if self.reader.readtext(img, detail=0):
                    used_ocr = True
                    logger.info("Image OCR triggered on %s", file_path)
            except Exception:
                logger.warning("OCR detection failed on %s", file_path)
            doc = self._convert(self.default_converter, proc)

        # Other formats
        else:
            doc = self._convert(self.default_converter, proc)
            used_ocr = False

        logger.debug("Doc structure for %s: %s", file_path, dir(doc))
        chunks = self.chunker.chunk(doc)
        return self._chunks_to_documents(chunks, file_path, doc, used_ocr)

    def _process_excel_with_pandas(self, file_path: str) -> List[Document]:
        docs: List[Document] = []
        try:
            sheets = pd.read_excel(file_path, sheet_name=None, header=None)
        except Exception as e:
            logger.error("Pandas fallback read failed for %s: %s", file_path, e)
            raise

        for sheet_name, df in sheets.items():
            df = df.dropna(how="all", axis=0).dropna(how="all", axis=1)
            if df.empty:
                continue

            header_row = df.iloc[0].fillna("").astype(str).tolist()
            data_rows = df.iloc[1:]
            lines = data_rows.astype(str).apply(lambda r: " | ".join(r), axis=1).tolist()

            buf, buf_len, idx = [], 0, 0
            for line in lines:
                ln = len(line) + 1
                if buf and buf_len + ln > self.chunk_size:
                    content = "\n".join(buf)
                    meta = {
                        "source": file_path,
                        "sheet": sheet_name,
                        "fallback": True,
                        "chunk_index": idx,
                        "type": "table",
                        "column_headers": header_row,
                    }
                    docs.append(Document(page_content=content,
                                         metadata=self._sanitize_metadata(meta)))
                    idx += 1
                    buf, buf_len = [], 0
                buf.append(line)
                buf_len += ln
            if buf:
                content = "\n".join(buf)
                meta = {
                    "source": file_path,
                    "sheet": sheet_name,
                    "fallback": True,
                    "chunk_index": idx,
                    "type": "table",
                    "column_headers": header_row,
                }
                docs.append(Document(page_content=content,
                                     metadata=self._sanitize_metadata(meta)))
        return docs

    def _convert(self, conv: DocumentConverter, path: str) -> Any:
        logger.info("Converting %s", path)
        try:
            out = conv.convert(path)
            return getattr(out, "document", out)
        except Exception as e:
            logger.error("Convert failed for %s: %s", path, e)
            raise

    def _chunks_to_documents(
        self,
        chunk_iter,
        file_path: str,
        docling_doc: Any,
        used_ocr: bool,
    ) -> List[Document]:
        docs: List[Document] = []
        current_section: Optional[str] = None
        captions = self._extract_vlm_captions(docling_doc)

        # build header lookup
        table_headers: Dict[Any, List[str]] = {}
        if hasattr(docling_doc, "pages"):
            for page in docling_doc.pages:
                for elem in getattr(page, "elements", []):
                    if getattr(elem, "type", "") == "table":
                        hdrs = [c.text for c in getattr(elem, "header_cells", [])]
                        table_headers[getattr(elem, "id", None)] = hdrs

        for chunk in chunk_iter:
            text = getattr(chunk, "text",
                           getattr(chunk, "content", str(chunk))).strip()
            if not text:
                continue

            if len(text) < 120 and text[-1] not in ".!?":
                current_section = text

            ctype = self._infer_chunk_type(text)
            wrapped = self._apply_wrapping(text, ctype)

            meta: Dict[str, Any] = {
                "reference": file_path,
                "file_name": os.path.basename(file_path),
                "type": ctype,
                "section": current_section,
                "ocr_used": used_ocr,
            }
            # inject headers for native tables
            if getattr(chunk, "element_type", "") == "table":
                elem_id = getattr(chunk, "element_id", None)
                if elem_id in table_headers:
                    meta["column_headers"] = table_headers[elem_id]

            if hasattr(chunk, "page"):
                p = getattr(chunk, "page")
                meta["page"] = p
                if p in captions:
                    meta["vlm_caption"] = captions[p]

            if hasattr(docling_doc, "language") and docling_doc.language:
                meta["language"] = str(docling_doc.language).lower()

            # smart split markdown tables
            lines = wrapped.splitlines()
            if (
                len(lines) >= 3 and
                lines[0].startswith("|") and
                re.match(r"^\|\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|$", lines[1])
            ):
                header, sep, *data = lines
                for i in range(0, len(data), self.max_table_rows):
                    block = data[i : i + self.max_table_rows]
                    content = "\n".join([header, sep] + block)
                    docs.append(Document(page_content=content,
                                         metadata=self._sanitize_metadata(meta.copy())))
                continue

            # normal split
            for sub in self._split_text(wrapped, self.chunk_size, ctype):
                docs.append(Document(page_content=sub,
                                     metadata=self._sanitize_metadata(meta.copy())))

        return docs

    def _extract_vlm_captions(self, doc: Any) -> Dict[int, str]:
        out: Dict[int, str] = {}
        try:
            for page in getattr(doc, "pages", []):
                for elem in getattr(page, "elements", []):
                    if getattr(elem, "type", "") == "image" and hasattr(elem, "caption"):
                        out[page.page_num] = elem.caption
        except Exception as e:
            logger.warning("VLM caption extract failed: %s", e)
        return out

    @staticmethod
    def _infer_chunk_type(text: str) -> str:
        t = text.strip()
        if not t:
            return "unknown"
        if len(t) < 120 and t[-1] not in ".!?":
            return "heading"
        if re.match(r"^(\-|\*|•|\d+\.)\s", t):
            return "list_item"
        if "\n" in t:
            letters = sum(c.isalpha() for c in t)
            symbols = sum((not c.isspace() and not c.isalnum()) for c in t)
            if symbols > letters:
                return "code"
        if "=" in t or re.search(r"[α-ω∞∑±≠≈]", t):
            return "formula"
        return "paragraph"

    @staticmethod
    def _apply_wrapping(text: str, ctype: str) -> str:
        if ctype == "code":
            return f"```\n{text}\n```"
        if ctype == "formula":
            return (
                f"$$\n{text}\n$$" if ("\n" in text or len(text) > 40)
                else f"${text}$"
            )
        return text

    @staticmethod
    def _split_text(text: str, max_size: int, ctype: str) -> List[str]:
        if len(text) <= max_size or ctype in ("code", "formula"):
            return [text]
        segments = re.split(r"(?<=[.!?。])\s+", text) or text.splitlines()
        chunks, buf = [], ""
        for seg in segments:
            if not seg.strip():
                continue
            if len(buf) + len(seg) + 1 <= max_size:
                buf += seg.strip() + " "
            else:
                if buf:
                    chunks.append(buf.strip())
                buf = seg.strip() + " "
        if buf:
            chunks.append(buf.strip())
        return chunks

    @property
    def supported_file_types(self) -> List[str]:
        return [
            "pdf", "docx", "pptx",
            "html", "htm", "xhtml",
            "md", "markdown", "txt", "csv", "asciidoc", "adoc",
            *self._image_exts,
            *self._excel_exts,
        ]

    def load_directory(self, directory: str) -> List[Document]:
        all_docs: List[Document] = []
        for root, _, files in os.walk(directory):
            for fname in files:
                ext = os.path.splitext(fname)[1].lower().lstrip(".")
                if ext in self.supported_file_types:
                    path = os.path.join(root, fname)
                    logger.info("Loading %s", path)
                    try:
                        docs = self.load_file(path)
                        for d in docs:
                            d.metadata["reference"] = path
                            d.metadata["file_name"] = fname
                        all_docs.extend(docs)
                    except Exception as e:
                        logger.error("Failed to load %s: %s", path, e)
                else:
                    logger.warning("Skipping unsupported: %s", fname)
        logger.info("Finished loading %d docs", len(all_docs))
        return all_docs
